"""Complete, clean data: typed values, the "Describe what you need" request, structured data, item pages,
failed-page retries, quality reports, run comparison and typed exports."""
import io
import json
import time
from datetime import date

import pytest
from cryptography.fernet import Fernet
from fastapi.testclient import TestClient

from app import config, db, extractor, fetcher, jobs, normalize, quality, recipe_builder, request_parser, \
    structured, vault

PUBLIC = "http://93.184.216.34"


# ---- typed values ------------------------------------------------------------------------

@pytest.mark.parametrize("text, amount, currency", [
    ("₹ 1,62,499", 162499, "INR"), ("Price: 150 INR/Piece", 150, "INR"), ("$1,299.99", 1299.99, "USD"),
    ("£51.77", 51.77, "GBP"), ("2.5 Lakh", 250000, ""), ("1.2 Cr", 12000000, ""), ("Rs. 12k", 12000, "INR"),
])
def test_money(text, amount, currency):
    got = normalize.parse_money(text)
    assert got[0] == amount and got[1] == currency


def test_ratings_numbers_dates():
    assert normalize.parse_rating("4.5 out of 5 stars") == 4.5 and normalize.parse_rating("Three") == 3
    assert normalize.parse_number("(151.6K)") == 151600
    assert normalize.parse_percent("-12%") == -12
    today = date(2026, 10, 23)
    assert normalize.parse_date("2d ago", today) == "2026-10-21"
    assert normalize.parse_date("Oct 23, 2026", today) == "2026-10-23"
    assert normalize.parse_date("23/10/2026", today) == "2026-10-23"
    assert normalize.parse_date("Featured", today) is None


def test_detect_and_normalize_columns():
    rows = [{"name": "A", "price": "₹ 450 / Piece", "rating": "4.5", "url": "https://x.test/a",
             "image": "https://x.test/a.jpg", "reviews": "(1,200)"},
            {"name": "B", "price": "₹ 1,250 / Piece", "rating": "4.1", "url": "https://x.test/b",
             "image": "https://x.test/b.jpg", "reviews": "(30)"}]
    types = normalize.detect_types(rows, list(rows[0]))
    assert types == {"name": "text", "price": "money", "rating": "rating", "url": "url", "image": "image",
                     "reviews": "number"}
    out = normalize.normalize_rows(rows, types)[0]
    assert out["price"] == 450 and out["price_currency"] == "INR" and out["price_unit"] == "piece"
    assert out["reviews"] == 1200 and out["rating"] == 4.5
    # A column named like a price but holding words stays text.
    assert normalize.detect_type("price", ["See options", "Ask for price"]) == "text"


# ---- request ---------------------------------------------------------------------------------

SHOES_REQUEST = {
    "columns": [{"name": "image", "meaning": "main picture", "type": "image", "required": True},
                {"name": "price", "meaning": "selling price", "type": "money", "required": True},
                {"name": "brand", "meaning": "brand; first word of the title", "type": "text", "required": True},
                {"name": "name", "meaning": "", "type": "text", "required": False}],
    "filters": [{"column": "price", "op": "lt", "value": "₹3,000"},
                {"column": "name", "op": "contains", "value": "men"}],
    "sort": {"column": "price", "desc": False}, "limit": 0, "auto": False, "details_hint": False,
    "question": "", "summary": "Image, price and brand; men's under ₹3,000; cheapest first.",
}


def test_understand_with_fake_ai(monkeypatch):
    monkeypatch.setattr(config, "GROQ_API_KEY", "k")
    monkeypatch.setattr(config, "AI_PROVIDER", "groq")
    seen = {}

    def fake_ask(system, prompt, schema, name):
        seen["prompt"], seen["schema"] = prompt, schema
        return SHOES_REQUEST
    monkeypatch.setattr(recipe_builder, "ask_json", fake_ask)
    req = request_parser.understand("I need image, price and brand. Men's shoes under ₹3,000, cheapest first",
                                    "https://shop.test/shoes")
    assert "₹3,000" in seen["prompt"] and seen["schema"]["required"]
    assert [c["name"] for c in req["columns"]] == ["image", "price", "brand", "name", "url"]
    assert req["columns"][-1]["auto_added"] and req["sort"] == {"column": "price", "desc": False}
    assert request_parser.describe(req) == ["price < ₹3,000", "name contains men", "sorted by price (low → high)"]


def test_understand_without_ai_lists_columns(monkeypatch):
    monkeypatch.setattr(config, "GROQ_API_KEY", "")
    monkeypatch.setattr(config, "ANTHROPIC_API_KEY", "")
    req = request_parser.understand("I need the image, price and brand of each shoe")
    assert [c["name"] for c in req["columns"]][:3] == ["image", "price", "brand"] and not req["auto"]


def test_filters_sort_limit_on_typed_values():
    req = request_parser.clean_request(SHOES_REQUEST)
    rows = [{"name": "Men's Runner", "price": 2499}, {"name": "Women's Runner", "price": 1999},
            {"name": "Men's Boot", "price": 4999}, {"name": "Men's Slide", "price": 899},
            {"name": "Men's Sandal", "price": ""}]
    types = {"price": "money", "name": "text"}
    matched = request_parser.filter_rows(rows, req, types)
    assert [r["name"] for r in matched] == ["Men's Runner", "Men's Slide"]     # not "Women's"
    assert [r["price"] for r in request_parser.sort_and_limit(matched, req)] == [899, 2499]
    top = request_parser.sort_and_limit(rows, {"sort": {"column": "price", "desc": True}, "limit": 2})
    assert [r["price"] for r in top] == [4999, 2499]


def test_recipe_prompt_carries_meanings():
    html = "<ul>" + "".join(f"<li class='c'><b>Nike Air Zoom Pegasus running shoe {i}</b><i>₹{i}99</i></li>" for i in range(5)) + "</ul>"
    req = request_parser.clean_request(SHOES_REQUEST)
    prompt = recipe_builder._make_prompt(recipe_builder.find_candidates(html), [], "https://shop.test/s",
                                         [c["name"] for c in req["columns"]], False, False, 11000, False, req)
    assert "brand (text): brand; first word of the title" in prompt and "price (money)" in prompt


# ---- structured data ---------------------------------------------------------------------------

LD_LIST = """<html><head><script type="application/ld+json">{"@context":"https://schema.org","@type":"ItemList",
"itemListElement":[%s]}</script></head><body><p>Showing 1-3 of 1,250 results</p></body></html>""" % ",".join(
    '{"@type":"ListItem","position":%d,"item":{"@type":"Product","name":"Shoe %d","brand":{"name":"Nike"},'
    '"url":"/p/%d","offers":{"price":"%d99","priceCurrency":"INR"}}}' % (i, i, i, i) for i in range(1, 4))


def test_jsonld_lists_and_json_recipe():
    lists = structured.find_lists(LD_LIST)
    assert lists and lists[0]["count"] == 3 and "offers.price" in lists[0]["keys"]
    recipe = extractor.normalize_recipe({"source": "json", "json_list": lists[0]["path"], "fields": [
        {"name": "name", "kind": "json", "selector": "name"},
        {"name": "brand", "kind": "json", "selector": "brand.name"},
        {"name": "price", "kind": "json", "selector": "offers.price"},
        {"name": "url", "kind": "json", "selector": "url"}]})
    rows = extractor.extract(LD_LIST, recipe, "https://shop.test/list")
    assert rows[0] == {"name": "Shoe 1", "brand": "Nike", "price": "199", "url": "https://shop.test/p/1"}
    assert quality.stated_total(LD_LIST) == 1250


def test_next_data_state_and_detail_jsonld():
    html = ('<script id="__NEXT_DATA__" type="application/json">{"props":{"pageProps":{"jobs":['
            + ",".join('{"title":"Dev %d","company":{"name":"Acme"},"salary":"%d LPA"}' % (i, i) for i in range(4))
            + ']}}}</script>')
    lst = structured.find_lists(html)[0]
    assert lst["path"] == "#__NEXT_DATA__:props.pageProps.jobs" and lst["keys"]["company.name"] == "Acme"
    assert structured.items_at(html, lst["path"])[3]["salary"] == "3 LPA"
    item = ('<script type="application/ld+json">{"@type":"JobPosting","title":"Dev","baseSalary":{"value":'
            '{"minValue":900000}},"hiringOrganization":{"name":"Acme"}}</script><main><h1>Dev</h1></main>')
    detail = extractor.normalize_recipe({"item_selector": "main", "fields": [
        {"name": "salary", "kind": "json", "selector": "baseSalary.value.minValue"},
        {"name": "title", "selector": "h1"}]})
    assert extractor.extract_detail(item, detail, "https://x.test/j/1") == {"salary": "900000", "title": "Dev"}


# ---- runs ------------------------------------------------------------------------------------

LIST = """<ul>%s</ul><a class="next" href="/list?page=2">next</a>""" % "".join(
    f"<li class='c'><b>Men's Shoe {i}</b><i>₹ {i},499</i><a href='/p/{i}'>x</a></li>" for i in range(1, 5))
LIST2 = """<ul><li class='c'><b>Men's Shoe 1</b><i>₹ 1,499</i><a href='/p/1'>x</a></li>
<li class='c'><b>Men's Shoe 9</b><i>₹ 999</i><a href='/p/9'>x</a></li></ul>"""
ITEM = "<main><h1>Shoe</h1><table><tr><th>Material</th><td>Mesh</td></tr></table></main>"
RECIPE = {"item_selector": "li.c", "fields": [{"name": "name", "selector": "b"}, {"name": "price", "selector": "i"},
                                              {"name": "url", "selector": "a", "attr": "href"}],
          "pagination": {"type": "next_link", "selector": "a.next"}}


@pytest.fixture
def fake_site(monkeypatch):
    vault.set_key_for_tests(Fernet.generate_key())
    db.init()
    state = {"fail": set(), "calls": []}

    def fetch_smart(url, recipe=None, max_rounds=0, use_cache=True):
        state["calls"].append(url)
        if url in state["fail"]:
            state["fail"].discard(url)          # fails once, works on the retry
            raise fetcher.FetchError(f"HTTP 503 for {url}")
        html = ITEM if "/p/" in url else LIST2 if "page=2" in url else LIST
        return fetcher.Page(url, url, 200, html, False)
    monkeypatch.setattr(fetcher, "fetch_smart", fetch_smart)
    monkeypatch.setattr(fetcher, "robots_allows", lambda url: True)
    yield state
    vault.set_key_for_tests(None)


def wait(job_id, timeout=5):
    for _ in range(int(timeout * 20)):
        job = db.get_job(job_id)
        if job["status"] not in ("queued", "running"):
            return job
        time.sleep(0.05)
    raise AssertionError("job did not finish")


def test_run_types_dedupes_filters_sorts_and_reports(fake_site):
    req = request_parser.clean_request({**SHOES_REQUEST, "limit": 2})
    job = wait(jobs.start(["https://shop.test/list"], {**RECIPE, "request": req}, max_pages=5))
    rows = db.get_records(job["id"])
    # 4 + 2 items, one duplicate (same link) -> 5 unique; shoes 3 and 4 cost 3,000 or more -> filtered
    # out; sorted by price and trimmed to 2.
    assert [r["price"] for r in rows] == [999, 1499]
    assert rows[0]["price_currency"] == "INR" and job["columns"]["price"] == "money"
    q = job["quality"]
    assert q["duplicates"] == 1 and q["excluded"] == 2 and q["trimmed"] == 1 and q["items_seen"] == 5
    assert q["columns"]["price"]["fill"] == 1.0 and q["pages_ok"] == 2


def test_failed_page_is_retried_at_the_end(fake_site):
    fake_site["fail"].add("https://shop.test/list?page=2")
    job = wait(jobs.start(["https://shop.test/list"], RECIPE, max_pages=5))
    assert job["status"] == "done" and job["rows"] == 5 and not job["quality"]["pages_failed"]
    assert fake_site["calls"].count("https://shop.test/list?page=2") == 2


def test_item_pages_fill_detail_columns(fake_site):
    recipe = {**RECIPE, "detail": {"url_field": "url", "item_selector": "main", "fields": [
        {"name": "material", "kind": "kv", "label_selector": "th", "value_selector": "td", "label": "Material"}]}}
    job = wait(jobs.start(["https://shop.test/list"], recipe, max_pages=1))
    rows = db.get_records(job["id"])
    assert len(rows) == 4 and all(r["material"] == "Mesh" for r in rows)
    assert job["quality"]["details"]["done"] == 4


def test_health_check_flags_columns_that_went_empty():
    prev = {"rows": 10, "columns": {"price": {"fill": 0.95}, "name": {"fill": 1.0}}}
    now = {"rows": 10, "columns": {"price": {"fill": 0.0}, "name": {"fill": 1.0}}}
    h = quality.health(now, prev)
    assert h["broken"] == ["price"] and quality.health(prev, prev) is None


# ---- API ------------------------------------------------------------------------------------

@pytest.fixture
def client(fake_site):
    from app.main import app
    return TestClient(app, base_url="http://localhost")


def test_preview_applies_the_request(client):
    req = request_parser.clean_request(SHOES_REQUEST)
    r = client.post("/api/preview", json={"url": PUBLIC + "/list", "recipe": {**RECIPE, "request": req}}).json()
    assert r["row_count"] == 4 and r["match_count"] == 2
    assert [row["price"] for row in r["rows"]] == [1499, 2499]
    assert r["column_types"]["price"] == "money" and r["recipe"]["column_types"]["price"] == "money"
    assert "price < ₹3,000" in r["request_labels"] and r["url_field"] == "url"


def test_compare_runs_and_typed_excel(client, fake_site):
    sid = client.post("/api/scrapes", json={"name": "Shoes", "urls": [PUBLIC + "/list"], "recipe": RECIPE,
                                            "max_pages": 1, "run": True}).json()
    first = sid["job_id"]
    wait(first)
    second = client.post(f"/api/scrapes/{sid['scrape_id']}/run", json={}).json()["job_id"]
    wait(second)
    diff = client.get(f"/api/jobs/{second}/compare").json()
    assert diff["previous_job"] == first and diff["counts"] == {"added": 0, "removed": 0, "changed": 0, "same": 4}
    x = client.get(f"/api/jobs/{second}/export?fmt=xlsx")
    import openpyxl
    ws = openpyxl.load_workbook(io.BytesIO(x.content)).active
    header = [c.value for c in ws[1]]
    assert isinstance(ws.cell(2, header.index("price") + 1).value, (int, float))


def test_understand_endpoint_without_ai(client, monkeypatch):
    monkeypatch.setattr(config, "GROQ_API_KEY", "")
    monkeypatch.setattr(config, "ANTHROPIC_API_KEY", "")
    r = client.post("/api/understand", json={"prompt": "title, company and salary"}).json()
    assert [c["name"] for c in r["columns"]] == ["title", "company", "salary", "url"]
